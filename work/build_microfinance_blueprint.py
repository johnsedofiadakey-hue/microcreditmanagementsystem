from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_ORIENT
from pathlib import Path
from datetime import date

OUT = Path('/Users/truth/Documents/Codex/2026-09-27/referenced-chatgpt-conversation-this-is-an/outputs/Microfinance_Susu_Internal_Technical_Blueprint.docx')

NAVY = '17365D'
BLUE = '1F5A8A'
PALE = 'EAF2F8'
PALE2 = 'F5F8FB'
GRAY = 'D9E1E8'
TEXT = '1F2933'
GREEN = 'EAF6EF'
AMBER = 'FFF3D6'
RED = 'FBE9E7'

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
normal.font.size = Pt(9.4)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(5.5)
normal.paragraph_format.line_spacing = 1.12
for name, size, before, after in [('Title', 26, 0, 14), ('Heading 1', 16, 14, 7), ('Heading 2', 12, 10, 5), ('Heading 3', 10.5, 8, 3)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name != 'Heading 3' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:ascii'), st.font.name)
    st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
# LibreOffice can inherit a decorative rule from the built-in Title style.
# The document title is intentionally typography-only.
title_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)

if 'Small Note' not in styles:
    note_style = styles.add_style('Small Note', WD_STYLE_TYPE.PARAGRAPH)
else:
    note_style = styles['Small Note']
note_style.base_style = styles['Normal']
note_style.font.size = Pt(8)
note_style.font.color.rgb = RGBColor.from_string('536270')

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), color)

def set_cell_border(cell, color=GRAY):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)

def set_cell_margins(cell, top=85, start=95, bottom=85, end=95):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    m = tcPr.first_child_found_in('w:tcMar')
    if m is None:
        m = OxmlElement('w:tcMar'); tcPr.append(m)
    for side, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = m.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side); m.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def add_text(cell, text, bold=False, size=8.3, color=TEXT, align=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.04
    if align is not None: p.alignment = align
    r = p.add_run(str(text)); r.bold = bold; r.font.name = 'Aptos'; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    return p

def table(headers, rows, widths=None, font=8.3, shade_rows=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, NAVY); set_cell_border(c, 'FFFFFF'); set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_text(c, h, True, font, 'FFFFFF')
        if widths: c.width = Inches(widths[i])
    trPr = t.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement('w:tblHeader'); header_repeat.set(qn('w:val'), 'true'); trPr.append(header_repeat)
    no_split = OxmlElement('w:cantSplit'); trPr.append(no_split)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        rowPr = t.rows[-1]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); rowPr.append(no_split)
        for i, val in enumerate(row):
            c = cells[i]
            shade(c, PALE2 if shade_rows and ridx % 2 else 'FFFFFF')
            set_cell_border(c); set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_text(c, val, False, font)
            if widths: c.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def p(text='', style=None, bold_lead=None):
    para = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = para.add_run(bold_lead); r.bold=True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para

def bullets(items):
    for item in items:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Inches(0.25)
        q.paragraph_format.first_line_indent = Inches(-0.16)
        q.paragraph_format.space_after = Pt(2)
        r = q.add_run('•  '); r.font.name = 'Aptos'; r.font.size = Pt(9.4)
        q.add_run(item)

def num(items):
    for index, item in enumerate(items, start=1):
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Inches(0.30)
        q.paragraph_format.first_line_indent = Inches(-0.20)
        q.paragraph_format.space_after = Pt(2)
        r = q.add_run(f'{index}.  '); r.font.name = 'Aptos'; r.font.size = Pt(9.4)
        q.add_run(item)

def tag(label, text, color):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(2); para.paragraph_format.space_after = Pt(5)
    run = para.add_run(label + ' '); run.bold = True; run.font.color.rgb = RGBColor.from_string(color)
    para.add_run(text)

def heading(text, level=1):
    para = doc.add_paragraph(text, style=f'Heading {level}')
    para.paragraph_format.space_after = Pt(11 if level == 1 else 7)
    return para

def page_break():
    doc.add_page_break()

def add_page_number(paragraph):
    paragraph.add_run('Page ')
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    paragraph._p.append(fldChar1); paragraph._p.append(instrText); paragraph._p.append(fldChar2)

def footer(section):
    f = section.footer
    para = f.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)
    r = para.add_run('Microfinance and Susu Internal Technical Blueprint | Confidential Internal Working Document')
    r.font.name='Aptos'; r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string('667788')

footer(sec)

# Cover
for _ in range(4): p('')
cover = p('Microfinance and Susu Management System', style='Title')
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = p('Internal Technical Blueprint', style='Heading 1'); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p('Architecture, controls, delivery plan and implementation decisions', style='Small Note').alignment = WD_ALIGN_PARAGRAPH.CENTER
for _ in range(7): p('')
table(['Document control', 'Value'], [
    ['Audience', 'Stormglide delivery team, client technical sponsor, finance lead, compliance lead and operations leadership'],
    ['Status', 'Internal planning blueprint - implementation is not authorized by this document alone'],
    ['Version', '1.0 | 27 September 2026'],
    ['Source basis', 'Microfinance and Susu Management System Project Blueprint (24 pages) plus the recorded KYC and fallback decisions'],
    ['Primary outcome', 'A build-ready, controlled architecture for a cash-recording financial operating system, beginning with Susu'],
], [1.45, 5.65], 8.8)
p('Prepared for internal technical planning. All regulatory statements are design assumptions to validate with the client, their legal counsel, Bank of Ghana requirements, NIA and contracted providers before production use.', style='Small Note')
page_break()

heading('Document control and reading guide', 1)
p('This document converts the client-facing project blueprint into an internal build plan. It preserves the intended operating model while making the boundaries, controls, responsibilities and unknowns explicit.')
table(['Marker', 'Meaning', 'Treatment'], [
    ['[Baseline]', 'Requirement or operating assumption already stated in the supplied project blueprint.', 'Retained unless a later decision supersedes it.'],
    ['[Agreed decision]', 'Direction stated in the referenced planning conversation.', 'Treat as current design direction; validate with the client during Phase 0.'],
    ['[Technical recommendation]', 'Stormglide architecture decision proposed to make the solution safe and buildable.', 'Requires Phase 0 acceptance before implementation.'],
    ['[Open decision]', 'A consequential item the client or regulated partner must decide.', 'Do not infer; record owner, deadline and evidence.'],
], [1.45, 3.25, 2.4], 8.2)
heading('Source reconciliation', 2)
table(['Topic', 'What the baseline says', 'Normalized technical position'], [
    ['Scope', 'One connected platform for Susu, microcredit, hire purchase, welfare, branches and about 200 paired field agents.', 'Retain. Build the common financial core once, then expose it through controlled product modules and channels.'],
    ['Money movement', 'The initial system records and tallies cash; it does not move money or require Mobile Money integration.', 'Retain for initial release. Any later payment collection or payout is a separate regulated integration programme.'],
    ['Identity', 'Ghana Card verification is used at registration; KYC is required.', 'Ghana Card and approved NIA IVSP verification are the primary identity route. Android phone biometrics only unlock the agent app; they do not supply a customer fingerprint to NIA.'],
    ['NHIS', 'Not defined as a primary KYC route.', 'Supporting captured card image and member-entered details only; never accepted as primary KYC or a substitute for Ghana Card verification.'],
    ['Roadmap count', 'Executive summary says 4 phases; delivery table lists Phase 0 through Phase 4.', 'Use Phase 0 (discovery and controls) plus four delivery phases. This is five stages, followed by ongoing support.'],
], [1.2, 2.7, 3.2], 7.8)
tag('[Critical boundary]', 'No client channel, Firebase product, agent device, SMS/USSD callback or portal browser may post ledger entries directly. All financial state changes go through the Financial Core API.', 'B22222')

heading('Executive architecture decision', 1)
p('The recommended design is a hybrid Firebase and PostgreSQL architecture on Google Cloud. Firebase is the delivery and trusted-device platform. PostgreSQL is the authoritative financial system of record. Cloud Run hosts the server-side financial and integration services. This maintains fast mobile and web delivery without asking a document database to enforce banking-grade transactions, locking and accounting invariants.')
table(['Layer', 'Authoritative responsibility', 'Primary technology'], [
    ['Channels', 'Present workflows and collect input; never decide balances or approvals.', 'Next.js PWA and public website; native Android app; USSD and SMS adapters'],
    ['Identity and app trust', 'Authenticate users, attest apps, register devices and deliver push configuration.', 'Firebase Authentication, App Check, FCM, Remote Config, Crashlytics, Analytics'],
    ['Financial Core', 'Validate commands, create double-entry journal, enforce idempotency, approval and branch rules.', 'TypeScript services on Cloud Run plus PostgreSQL'],
    ['Evidence and documents', 'Store encrypted photographs, agreements and KYC documents with immutable references.', 'Firebase Storage / Google Cloud Storage with server-only signed access'],
    ['Integration edge', 'NIA IVSP, USSD, SMS, credit bureau, bank statements, MDM and optional future payments.', 'Cloud Run adapter services, Pub/Sub and provider-specific queues'],
], [1.35, 3.25, 2.45], 8.0)

heading('1 System context goals and boundaries', 1)
tag('[Baseline]', 'The business serves community members through Susu savings, micro credit, hire purchase and welfare. Agents work in pairs, branches receive and pay cash, and management needs same-day visibility and cash controls.', BLUE)
tag('[Agreed decision]', 'The system must have a controlled fallback when biometric and external KYC verification fails; a failed scan must not become either an automatic rejection or an easy staff override.', BLUE)
heading('1.1 Goals', 2)
bullets([
    'Trace each collection from member through agent pair, branch handover, branch close and bank reconciliation, with a durable receipt and evidence trail.',
    'Make the member record, financial position and requests consistent across agent app, member PWA, USSD, SMS, branch portal and reports after confirmed sync.',
    'Support financial products only through configured business rules, limit approvals and append-only journal entries.',
    'Operate through intermittent connectivity without double-posting, losing receipts or silently changing the official ledger.',
    'Create regulatory-ready evidence, auditable decision history and operational controls without pretending the application itself is a licence or compliance certification.'
])
heading('1.2 Explicit non-goals for the first production scope', 2)
bullets([
    'No Mobile Money collection, disbursement or wallet custody. USSD and PWA create view or request actions only; cashiers and branches perform the approved cash payout.',
    'No consumer app may read raw NIA biometrics or use a phone fingerprint sensor as a Ghana Card fingerprint scanner.',
    'No editable balances, deleted transactions, staff self-approval, direct Firestore ledger writes or client-held provider secrets.',
    'No automated credit decision that bypasses a configured policy, human review threshold or adverse-decision reason where those are required.'
])

heading('2 Target system and hosting topology', 1)
p('One Google Cloud organization should host separate Firebase projects and GCP projects for development, training/UAT and production. Production services run in a dedicated production project, with least-privilege service accounts and a separately governed billing account.')
table(['Zone', 'Components', 'Trust boundary and notes'], [
    ['Public edge', 'Firebase Hosting for Next.js public site and member PWA; CDN; Cloud Armor if the selected hosting topology requires it.', 'Only static/SSR delivery and BFF calls. No private database access from browser code.'],
    ['Managed clients', 'Android agent app, browser PWA, admin/branch portal, USSD and SMS provider callbacks.', 'Firebase App Check validates app origin where supported. Device binding and role policy remain server-enforced.'],
    ['Application services', 'Cloud Run: Financial Core API, Sync API, USSD gateway, KYC gateway, document service, reporting worker, notification dispatcher.', 'Private service-to-service identity. Egress and provider credentials limited by service.'],
    ['Financial data plane', 'Managed PostgreSQL with HA, PITR, encrypted backups, read replica/reporting path where justified.', 'Only Financial Core and authorised worker service accounts connect. No mobile/browser connection strings.'],
    ['Async and evidence plane', 'Pub/Sub, Cloud Tasks, Cloud Scheduler, Firebase Storage / GCS, audit event sink, monitoring logs.', 'Outbox events are published after database commit. Documents are referenced by immutable IDs and hashes.'],
], [1.4, 3.1, 2.55], 7.8)
heading('2.1 Request and event flow', 2)
p('Member or staff command -> Firebase-authenticated/App Check-attested channel -> Cloud Run API -> policy and idempotency checks -> PostgreSQL transaction and journal -> outbox event -> notification, reporting, fraud or reconciliation worker. The committed journal is the only source for balance projection.')

heading('3 Exact recommended stack and responsibility boundary', 1)
table(['Capability', 'Recommended stack', 'Why this is the chosen boundary'], [
    ['Public site and member PWA', 'Next.js 15+ / React / TypeScript, Firebase Hosting with SSR through Cloud Run where needed.', 'Fast public delivery and progressive member experience; works alongside USSD rather than replacing it.'],
    ['Admin and branch portal', 'Next.js / TypeScript / React Query / accessible component system on Firebase Hosting.', 'Role-aware operational interface; browser calls only server APIs.'],
    ['Agent Android app', 'Native Android with Kotlin, Jetpack Compose, Room/SQLCipher local store, WorkManager sync, Android Enterprise managed configuration.', 'Best control over offline data, device policy, background sync, hardware and managed distribution.'],
    ['Financial services', 'TypeScript with NestJS or Fastify on Cloud Run; OpenAPI contracts; Cloud Tasks and Pub/Sub.', 'Typed, independently deployable commands and adapter services with managed scaling.'],
    ['Financial core database', 'PostgreSQL 16 managed on Cloud SQL or AlloyDB depending scale and availability requirement.', 'ACID transactions, constraints, locking, ledger journal, reporting SQL, immutable history and safe reconciliation.'],
    ['Operational cache', 'Redis / Memorystore only for ephemeral rate limits, distributed locks and session helpers.', 'Never store the financial source of truth or approval state only in cache.'],
    ['Documents', 'Firebase Storage backed by GCS, CMEK where the selected data classification requires it.', 'Fine-grained storage policy and signed server-mediated access.'],
    ['Observability', 'Cloud Logging/Monitoring/Trace, Error Reporting, Firebase Crashlytics and Analytics, Sentry optional for cross-channel product errors.', 'Separates server operational evidence from client crash and usage signals.'],
    ['Infrastructure and delivery', 'Terraform, GitHub Actions, Artifact Registry, Cloud Build optional, Firebase CLI for Hosting only.', 'Repeatable reviewable changes and narrow deployments.'],
], [1.45, 2.45, 3.15], 7.55)
heading('3.1 Firebase must handle', 2)
bullets([
    'Firebase Authentication identity lifecycle, token issuance and basic claims; App Check; Firebase Hosting; FCM; Crashlytics; Analytics; Remote Config; Storage delivery policy; emulator-backed local development.',
    'Firestore only for non-financial, transient or replicated read models such as public content drafts, app remote configuration mirrors, in-app announcements and temporary operational presence if a relational record is not required.',
    'Cloud Functions only for Firebase-native triggers that do not create financial state, such as image post-processing, FCM fan-out helper calls or Analytics-to-notification enrichment. Use Cloud Run for all financial or regulated integrations.'
])
heading('3.2 Firebase must not handle', 2)
bullets([
    'The general ledger, member account balance, loan schedule authority, reconciliation match authority, approval chain, cash close, fraud decision evidence or immutable audit log.',
    'Raw NIA biometric data, card images as public client-accessible objects, provider secrets, private keys, database credentials or logic that relies solely on Firestore security rules for financial correctness.',
    'Direct client-generated financial writes, scheduled mass financial postings or a substitute for double-entry accounting and database transaction isolation.'
])

heading('4 Channels and product surfaces', 1)
heading('4.1 Member customer PWA', 2)
tag('[Technical recommendation]', 'The PWA is a richer, optional self-service companion. It must offer the same authoritative view and request behaviours as USSD but is not required for access to core services.', BLUE)
table(['Surface', 'Initial capabilities', 'Controls'], [
    ['Member PWA', 'Profile review; account and transaction read model; statements; request withdrawal, loan, hire purchase or assistance; case status; branch contact; device/session management.', 'Firebase Auth phone or passkey-assisted sign-in combined with PIN step-up for sensitive actions; server risk checks; masked information; SMS confirmations.'],
    ['Public website', 'Brand, locations, services, calculators labelled illustrative, KYC process explanation, interest form, accessibility and privacy content.', 'No financial account data. Rate-limit forms, CAPTCHA/App Check alternative, consent capture and a route to the correct branch.'],
    ['Admin/branch portal', 'Work queues, branch close, reconciliation, approvals, cases, reports, product configuration and audit review.', 'Workforce MFA, managed browser/device expectation, short sessions, re-authentication for high-risk approval and policy-backed data scoping.'],
    ['Agent Android', 'Offline registration intake, collections, route, cash handover, request initiation, supervisor checks and queued evidence capture.', 'Managed device, device binding, local encryption, biometric local unlock, server-issued session, sync proof and remote revoke.'],
    ['USSD/SMS', 'Balance and recent activity views, low-risk requests, status, PIN changes under safeguards and outbound receipts/alerts.', 'PIN attempt limits, telco/MSISDN risk flags, non-sensitive screen limits, callback signatures, replay protection and SMS delivery state.'],
], [1.25, 3.35, 2.45], 7.8)
heading('4.2 USSD and SMS behaviour', 2)
p('Use a provider-neutral USSDGateway interface with adapters for mNotify, NALO, Arkesel and Africa\'s Talking. Start provider selection with commercial and coverage validation, not hard-coded provider assumptions. The interface exposes startSession, continueSession and endSession; all session state is short-lived and non-financial. The Financial Core receives a signed, idempotent command from the gateway only after session and PIN policy validation.')
table(['Action', 'USSD/PWA allowed?', 'Financial result'], [
    ['Check balance/recent activity', 'Yes', 'Read-only projection with a "last confirmed" timestamp.'],
    ['Request withdrawal', 'Yes', 'Creates a request and optional reservation under configured policy. No balance debit until approved payout posts.'],
    ['Request loan/hire purchase/welfare support', 'Yes', 'Creates a case/application. No approval, disbursement or value-score override in the channel.'],
    ['Pay, transfer or disburse money', 'No in initial scope', 'Deferred to a separately governed payment programme.'],
    ['PIN reset', 'No self-service reset over USSD alone', 'Agent/branch-assisted recovery with KYC and fraud controls.'],
], [2.15, 1.65, 3.25], 8.0)

heading('5 Identity KYC account recovery and device trust', 1)
heading('5.1 Ghana Card and NIA IVSP primary KYC', 2)
tag('[Agreed decision]', 'Ghana Card is the primary identity source. NHIS is only supporting captured ID data/photo and cannot approve primary KYC.', BLUE)
table(['Stage', 'Design', 'State/result'], [
    ['Capture', 'Agent captures Ghana Card fields, consent, member photo and required supporting records. The app validates completeness offline but cannot declare KYC complete.', 'KYC_DRAFT'],
    ['Online verification', 'Cloud Run KYC gateway sends only the approved NIA/IVSP request through a contracted provider. Keep provider response reference, method, outcome, timestamp and policy version.', 'KYC_VERIFIED or KYC_RETRY_REQUIRED'],
    ['Biometric route', 'Use approved IVSP face/liveness route or a NIA-supported fingerprint reader/MECO device. Phone fingerprint scanner only authenticates the device owner locally.', 'Method recorded; no raw biometric template stored by application'],
    ['Controlled exception', 'Retry conditions first: lighting, capture quality, alternate supported method, IVSP escalation. If still unresolved, create a compliance case.', 'KYC_EXCEPTION and restricted account'],
    ['Manual review', 'Compliance receives evidence, provider case ID and reason. They may request NIA/IVSP escalation or rectification; no generic override button.', 'RESTRICTED_KYC, then VERIFIED or REJECTED/EXPIRED per policy'],
], [1.45, 4.35, 1.25], 7.6)
p('Recommended restricted state: the member profile can exist and staff can preserve the evidence, but the product configuration enforces the client-approved post-no-debit / restricted rules until resolution. The exact 90-day rectification treatment and what activity is permitted must be confirmed against the applicable NIA and Bank of Ghana guidance before launch.')
heading('5.2 Account recovery', 2)
bullets([
    'Member recovery: verify Ghana Card and run a fraud check before changing phone number or resetting PIN. A recent phone change plus withdrawal request must trigger a cooling period or dual review policy.',
    'Agent recovery: supervisor-assisted device re-enrolment; revoke old device token; verify employee identity, MDM posture and handover of un-synced encrypted queue when recoverable.',
    'Staff recovery: helpdesk ticket plus workforce MFA recovery and role re-approval. Privileged role recovery requires a separate approver and audit review.',
    'Never use SMS OTP alone as the sole recovery factor for an account with financial requests. Treat SIM swap and port-out signals as a risk input where available.'
])
heading('5.3 Device binding', 2)
p('The agent app receives a device registration certificate only after Android Enterprise enrolment, managed app configuration and employee binding. The app stores a hardware-backed key where available, attaches attestation evidence and a device ID to each command, and rejects action when the device is revoked, non-compliant or outside the allowed application version grace window. Firebase App Check is an additional signal, not the device-binding system by itself.')

heading('6 Authoritative financial core', 1)
tag('[Baseline]', 'The supplied blueprint requires a double-entry, append-only ledger, unique transaction IDs, server time as official time and reversals rather than deletions.', BLUE)
heading('6.1 Financial data model', 2)
table(['Domain record', 'Authoritative table group', 'Core invariants'], [
    ['Identity and organisation', 'members, member_identifiers, branches, wings, towns, users, agents, pairs, devices, roles', 'Historic assignment effective dates; no hard-delete for used identities; branch scope is explicit.'],
    ['Products and contracts', 'susu_products/accounts, loan_products/contracts, hp_catalog/contracts, welfare_memberships/claims', 'Versioned terms; contract captures the product rule version accepted at origination.'],
    ['Commands and evidence', 'commands, idempotency_keys, documents, cases, approvals, audit_events', 'Every state change links to authenticated actor, device/channel, correlation ID and evidence.'],
    ['Accounting', 'journal_batches, journal_lines, account_balances projection, chart_of_accounts, posting_rules', 'Each committed batch balances debit equals credit; posted lines immutable; reversal references original batch.'],
    ['Cash and reconciliation', 'collections, custody_events, handovers, branch_closes, bank_deposits, reconciliation_matches', 'Amounts cannot advance to a later custody stage without prior valid state; mismatch remains visible.'],
], [1.5, 3.1, 2.45], 7.8)
heading('6.2 Ledger posting rule', 2)
p('A business command never writes an account balance. It validates policy and creates a journal batch in one PostgreSQL transaction. A balance projection updates from the journal in the same transaction or a transactionally consistent worker. An idempotency key composed from channel, device/client command ID and intent prevents duplicate sync or webhook replay from posting twice. The client-issued UUID is evidence; the server assigns official posting ID and time.')
table(['Example business event', 'Illustrative accounting effect', 'Control'], [
    ['Cash Susu collection confirmed', 'Debit agent cash custody; credit member Susu liability.', 'Receipt only after commit; pair and route evidence attached.'],
    ['Agent handover accepted', 'Debit branch cash custody; credit agent cash custody.', 'Both sides confirm counted amount; variance opens a case.'],
    ['Cash withdrawal paid', 'Debit member Susu liability; credit branch cash custody.', 'Maker-checker request and cashier payout confirmation required.'],
    ['Loan disbursement paid', 'Debit loan receivable; credit branch cash custody.', 'Terms and approval chain frozen before payout.'],
    ['Correction', 'Post equal and opposite reversal batch, then corrected new batch if valid.', 'Original remains immutable; reason and approver required.'],
], [1.65, 3.25, 2.15], 7.7)
heading('6.3 Financial database controls', 2)
bullets([
    'Use database constraints, foreign keys, numeric(18,2) currency fields or integer minor-units consistently, posting-period locks and stored procedures/service transaction boundaries for journal posting.',
    'Use SERIALIZABLE or carefully designed transaction locking for high-risk reserve/payout flows; test contention and deadlock retry logic explicitly.',
    'Keep separation between application ledger and general-ledger export. The client accountant owns the final chart of accounts mapping and statutory accounting treatment.',
    'Use append-only database audit logs plus application audit events; database administrators cannot silently alter financial history without generating a security event and recovery evidence.'
])

heading('7 Offline first sync and field cash operations', 1)
heading('7.1 Offline sync contract', 2)
table(['Requirement', 'Design decision', 'Failure handling'], [
    ['Durable local records', 'Room database encrypted using SQLCipher or Android Keystore-wrapped database key. Persist command queue before UI success.', 'If storage/encryption fails, block collection and show a clear recoverable error; never keep an in-memory-only receipt.'],
    ['Exactly-once financial effect', 'At-least-once delivery with server idempotency. Each command has device UUID, sequence number, intended amount, member and evidence hash.', 'Duplicate requests return original server result, not a second posting.'],
    ['Ordering', 'Per-device monotonic sequence and dependency graph; server does not trust client clock.', 'Out-of-order commands remain pending or are rejected with explicit reason, not silently reordered.'],
    ['Conflicts', 'Financial commands are immutable; profile/product reads use version stamps. Agent cannot merge two financial outcomes locally.', 'Create a conflict/case or fetch authoritative state; supervisor resolves only through auditable correction.'],
    ['Receipts', 'Display provisional local receipt with "pending sync"; issue authoritative SMS/receipt after server commit.', 'If SMS is delayed, preserve commit receipt and retry notification through outbox.'],
    ['Remote revoke', 'Stop new commands when device is revoked after a limited offline grace policy.', 'Allow secure sync of already-queued signed commands only if policy permits; flag for review.'],
], [1.45, 3.85, 1.75], 7.5)
heading('7.2 Daily cash lifecycle', 2)
p('Collection -> agent pair cash position -> confirmed handover -> branch cashier custody -> branch count and close -> bank deposit -> bank statement match. Each stage uses its own event, counted amount, people, time, evidence and variance. The system must report both member liability and physical custody; they are related but not the same number.')
table(['Operational checkpoint', 'Owner', 'Required evidence'], [
    ['Collection', 'Agent pair', 'Member, amount, product, geo/route signal, time, device, provisional/confirmed receipt state.'],
    ['Handover', 'Pair and branch cashier', 'Expected vs counted amount, both acknowledgements, cash bag/receipt reference, variance reason.'],
    ['Branch close', 'Branch manager and cashier', 'Cash count, float, pending payouts, open variances, signed close.'],
    ['Deposit and reconciliation', 'Finance', 'Bank deposit reference, statement import line, match status, reviewer and unresolved exceptions.'],
], [1.7, 1.55, 3.8], 8.0)

heading('8 Product modules and maker checker', 1)
table(['Module', 'Build behaviour', 'Key control'], [
    ['Susu', 'Daily/weekly/monthly packages, multiple accounts, catch-up rules, contribution schedule, reserve/request/payout and statement.', 'Configured commission and withdrawal rules are versioned. Available balance considers approved reservations.'],
    ['Loans', 'Eligibility, credit review, guarantors/lien, approved terms, schedule, collections, PAR 30/60/90 and recovery cases.', 'Member Value Score is explainable, versioned and advisory/threshold-driven until client approves its policy.'],
    ['Hire purchase', 'Catalog, value bands, supplier, deposit, contract, delivery evidence, instalments and recovery workflow.', 'Supplier payment and member contract are separate approvals; delivery proof required before certain postings.'],
    ['Welfare', 'Membership, contributions, claims, evidence, policy-based decision and payout workflow.', 'Claim conflict checks and defined approver limits; payout remains a cash custody event.'],
    ['Cases and disputes', 'Member complaint, receipt dispute, KYC exception, cash variance, fraud alert, data correction and recovery case.', 'No direct balance edit. Resolution posts a linked adjustment/reversal only after evidence and approval.'],
], [1.15, 3.55, 2.35], 7.7)
heading('8.1 Maker-checker policy', 2)
table(['Action', 'Maker', 'Checker / rule'], [
    ['Withdrawal request', 'Agent, member channel or cashier intake', 'Branch approver within configured amount limit; cashier must be different actor for payout.'],
    ['Loan / hire purchase approval', 'Credit officer', 'Different authorised approver, tiered by amount and policy exception.'],
    ['Reversal / correction', 'Operations or finance maker', 'Independent finance/compliance checker; original transaction must be linked.'],
    ['Write-off / lien release / welfare payout', 'Responsible product officer', 'Independent approver at configured authority level; compliance evidence where required.'],
    ['Role/limit change', 'Identity administrator', 'Separate security or business owner approval; access review retains both records.'],
], [2.0, 2.05, 3.0], 7.8)

heading('9 APIs integrations and asynchronous processing', 1)
p('Use OpenAPI 3.1 contracts for every channel-to-core interface. Each request carries correlation ID, authenticated principal, channel, client command ID and idempotency key. Separate command endpoints from reporting/read endpoints. Publish a versioned domain event after successful commit using a transactional outbox pattern.')
table(['Provider / interface', 'Integration design', 'Production gate'], [
    ['NIA IVSP / licensed verifier', 'Cloud Run KYC gateway; provider adapter; encrypted request/response metadata; case correlation.', 'Contract, lawful purpose/consent language, sandbox or certification, error catalogue and data retention schedule.'],
    ['USSD gateway', 'Provider-neutral adapter; signed callbacks; per-session state in Redis with short TTL; core API command calls.', 'Network coverage, shared/dedicated code, SLA, callback verification and peak-session test.'],
    ['SMS provider', 'Outbound notification service; templates, delivery receipts, retry and cost guardrails.', 'Sender registration, delivery-report webhook, opt-out/legal wording and cost ceiling.'],
    ['Credit bureau', 'Batch/file/export adapter first unless a compliant API is contracted; immutable submission records and acknowledgement tracking.', 'Confirm licence obligation, file format, transmission method, correction process and 72-hour expectation applicability.'],
    ['Bank statement', 'Read-only file import initially (CSV/OFX/camt where possible), parser quarantine, matching workbench.', 'Bank format samples, period close policy, reconciliation owner and exception SLA.'],
    ['MDM / Android Enterprise', 'Managed Google Play private app distribution, compliance signal and remote lock/wipe integration.', 'Enterprise enrollment, company ownership policy, replacement-device procedure.'],
    ['Future payments', 'Separate payment orchestration and settlement/reversal service; not connected to ledger directly.', 'Regulatory/legal review, PSP contract, webhook security, settlement reconciliation, limits and fraud model.'],
], [1.55, 3.7, 1.8], 7.35)
heading('9.1 External failure design', 2)
bullets([
    'Treat provider callbacks as untrusted until signature, allow-list, timestamp/replay window and idempotency checks pass.',
    'Use Cloud Tasks for retryable outbound work, dead-letter queues for exhausted jobs and an operations queue with named owner. Never retry a financial command blindly without idempotency evidence.',
    'If NIA/USSD/SMS is unavailable, preserve local/queued intake or request state, display the limitation honestly and escalate according to policy. Do not invent a positive verification or delivery result.',
    'Store provider correlation IDs, request hashes and redacted error classes, not raw secrets or unnecessary biometric payloads.'
])

heading('10 Security privacy and audit', 1)
heading('10.1 Security design', 2)
table(['Control domain', 'Required baseline'], [
    ['Authentication', 'Firebase Authentication with staff MFA; custom claims are coarse routing only. Server evaluates active role, branch, limit and separation-of-duties policy on every command.'],
    ['Authorization', 'PostgreSQL policy service and service-layer checks; deny by default; branch/wing scope, product scope and effective dates. Quarterly access recertification.'],
    ['Secrets', 'Google Secret Manager; service identity access only; no secrets in Firebase config, app binaries, source control, analytics, tickets or logs. Rotate on personnel/provider change.'],
    ['Encryption', 'TLS 1.2+ in transit; database/storage encryption at rest; field-level encryption/tokenisation for highly sensitive identifiers; Android Keystore and encrypted local database.'],
    ['Network', 'Cloud SQL private IP or connector; no public database; Cloud Run ingress limited per service; provider egress allow-list where feasible; WAF/rate limits on public endpoints.'],
    ['Audit', 'Append-only application audit events with actor, before/after logical state, reason, correlation ID, source IP/device, time and linked approval. Export immutable security/audit logs to protected retention storage.'],
    ['Privacy', 'Data minimisation, explicit purpose and consent record, access/download/correction workflow, retention jobs, subject-request case handling and staff viewing reason for KYC images.'],
], [1.45, 5.55], 8.0)
heading('10.2 Data classification and retention assumptions', 2)
table(['Data class', 'Examples', 'Proposed retention / handling'], [
    ['Financial and audit', 'Journal, receipts, approvals, contracts, reconciliation and fraud cases.', 'At least 5 years after relationship/transaction unless applicable law or regulator requires longer; legal hold overrides deletion.'],
    ['KYC and identity', 'Ghana Card fields, verification result/reference, member image and supporting ID documents.', 'Minimum necessary period under agreed KYC/AML policy; encrypted, tightly scoped viewing and documented deletion review.'],
    ['NHIS support data', 'Captured card photo and entered supporting details.', 'Only if business purpose documented; not used as primary KYC; shorter retention if it has no continuing operational/legal purpose.'],
    ['Operational telemetry', 'Crash logs, app version, device health, USSD session diagnostics.', 'Minimise PII; separate retention by troubleshooting value, typically shorter than financial records.'],
], [1.35, 2.85, 2.8], 7.8)
p('These are system design assumptions, not a legal retention schedule. The client must adopt an approved data-retention policy and confirm the regulator, AML/FIC, privacy and contractual obligations before live data migration.')

heading('11 Backup disaster recovery and service levels', 1)
table(['Asset', 'Backup / recovery design', 'Target to validate'], [
    ['PostgreSQL financial core', 'Regional HA, automated daily backups, point-in-time recovery, encrypted off-project backup copy and monthly restore rehearsal.', 'Proposed RPO 15 minutes; proposed RTO 4 hours for core command service.'],
    ['Documents', 'Versioning, retention policy, cross-region replication only after data-residency assessment, restore test against sample case.', 'Proposed RPO 24 hours for non-financial evidence; tighter where provider evidence requires it.'],
    ['Configuration and infrastructure', 'Terraform state protection, versioned deployment manifests, Firestore/config export where used, Secrets inventory not secret value export.', 'Rebuild a clean environment from code and approved secret rotation process.'],
    ['Android offline queue', 'Encrypted device-local queue plus normal sync; no cloud backup of application private financial store.', 'Data loss risk accepted only for un-synced records after defined device-loss policy; field process mitigates with timely sync.'],
], [1.65, 3.8, 1.55], 7.5)
heading('11.1 Incident model and proposed operating SLOs', 2)
table(['Service', 'Proposed target', 'Measurement'], [
    ['Financial Core API', '99.9% monthly availability excluding approved maintenance; no accepted command without durable commit response.', 'Synthetic command/read probes, error rate, latency and failed posting metrics.'],
    ['Agent sync', '95% of healthy queued commands confirmed within 15 minutes of usable network return.', 'Queue age, retry count, device health and sync success dashboard.'],
    ['SMS receipt dispatch', '99% accepted by provider within 5 minutes of posted event; delivery is separately reported.', 'Outbox lag and provider delivery receipts.'],
    ['High severity incident', 'Acknowledge in 30 minutes during agreed support window; preserve evidence before remediation.', 'Incident ticket timeline, audit export and corrective-action review.'],
], [1.55, 3.5, 1.95], 7.8)

heading('12 Observability fraud controls and cases', 1)
table(['Signal', 'Detection logic', 'Response'], [
    ['Cash variance', 'Expected collection/handover/close/deposit mismatch above policy threshold.', 'Block branch close escalation as policy requires; create variance case; notify manager and audit.'],
    ['Collection anomaly', 'Drop in pair performance against baseline; repeated missed-member patterns; unusual time/geo distance.', 'Flag for supervisor field audit; do not automatically accuse or alter balance.'],
    ['Account takeover', 'Phone change, SIM risk input, new device or failed PIN sequence followed by withdrawal request.', 'Cooling period, step-up KYC and independent approval.'],
    ['Staff misuse', 'High reversal count, self-approval attempt, unusual after-hours action or permission escalation.', 'Deny self-approval, create security case and require compliance review.'],
    ['KYC failure', 'Repeated no-match, device quality issue or provider outage.', 'Retry/escalate ladder; restricted KYC state, not bypass.'],
], [1.5, 3.65, 1.85], 7.7)
p('A case is a first-class record with classification, severity, owner, SLA, evidence, communications, resolution and linked financial/KYC/audit IDs. It is not an informal comment thread. Closing a dispute requires an outcome code and, when money changes, a linked approved journal reversal or adjustment.')

heading('13 Environments CI CD and testing', 1)
heading('13.1 Environments', 2)
table(['Environment', 'Purpose', 'Data and controls'], [
    ['Development', 'Feature work, local emulators and contract testing.', 'Synthetic data only; disposable Firebase/GCP resources; no provider production credentials.'],
    ['Training / UAT', 'Agent and branch rehearsal; acceptance testing and migration dry runs.', 'Synthetic or irreversibly masked data; managed training devices; isolated project and provider sandbox/shared code.'],
    ['Production', 'Live member operations.', 'Real data only; change control; restricted break-glass access; separate billing, secrets, backups and log retention.'],
], [1.45, 2.6, 2.95], 8.0)
heading('13.2 Delivery pipeline', 2)
num([
    'Pull request: lint, type check, unit tests, dependency/security scan, API contract checks, database migration review and infrastructure plan review.',
    'Ephemeral or shared test deployment: integration tests with Firebase Emulator Suite, PostgreSQL test database, provider mocks and Android instrumentation tests.',
    'UAT promotion: immutable build artifact, signed Android package, release notes, migration rehearsal and acceptance checklist.',
    'Production deployment: approval gate; run additive/backwards-compatible migrations first; deploy Cloud Run revision; deploy Firebase Hosting narrowly; verify synthetic flows and a distinctive version marker.',
    'Rollback: route traffic to last healthy Cloud Run revision and previous Hosting release where safe. Do not roll back a committed financial migration or journal; use forward correction and an incident-controlled remediation plan.'
])
heading('13.3 Test strategy', 2)
table(['Test class', 'Evidence required before pilot'], [
    ['Financial correctness', 'Double-entry property tests, invariant tests, idempotency/replay tests, posting-period and reversal tests, concurrency/lock tests.'],
    ['Offline', 'Airplane mode, flaky network, killed process, clock shift, duplicate queue, storage pressure, device replacement and two-agent conflict simulation.'],
    ['Security', 'Role/branch escape tests, App Check behaviour, device revoke, secrets scan, API abuse/rate limit, webhook forgery and pen test before scale-out.'],
    ['Integration', 'Provider sandbox or certified test: NIA error taxonomy, SMS receipts, USSD menu/session expiry, bank import, MDM enrollment and FCM.'],
    ['Operational acceptance', 'One full day: collection to handover to close to bank reconciliation, then a dispute/reversal scenario and recovery drill.'],
], [1.65, 5.35], 8.0)

heading('14 Migration deployment and rollout', 1)
heading('14.1 Migration protocol', 2)
num([
    'Inventory paper cards, spreadsheets, branch cash records and existing member identifiers; produce a data-quality report before import.',
    'Map and cleanse into staging tables. Resolve duplicate people, missing identifiers and contradictory balances outside the production ledger.',
    'Import only signed opening positions with source references, branch approval and a balanced opening journal. Never fabricate detailed historic transactions if only opening balances are reliable.',
    'Send opening-balance confirmation through approved member channels and open time-bounded disputes before each branch goes live.',
    'Rehearse import, reconciliation and rollback in UAT. Production import is a controlled cutover with a before/after total signed by the branch manager and finance owner.'
])
heading('14.2 Normalized delivery stages', 2)
table(['Stage', 'Timing guide', 'Deliverable / release gate'], [
    ['Phase 0 - Discovery and controls', 'Month 1', 'Signed business rules, licence/category assumptions, data inventory, provider due diligence, threat model, acceptance criteria and pilot readiness plan.'],
    ['Phase 1 - Core Susu pilot', 'Months 2-3', 'Member/KYC intake, offline agent collection, ledger, branch handover/close, SMS receipt, public site v1 and one-branch pilot with about 10 pairs.'],
    ['Phase 2 - Member access and scale-out', 'Months 4-5', 'USSD/PWA view and request flows, withdrawal process, fraud alerts, commissions, branch waves, support operations and 200-agent readiness.'],
    ['Phase 3 - Credit and hire purchase', 'Months 6-7', 'Validated score policy, loans, schedules, bureau reporting/export, catalog/suppliers, delivery proof and portfolio controls.'],
    ['Phase 4 - Welfare reporting and hardening', 'Months 8-9', 'Welfare, regulatory/reporting pack, advanced dashboards, resilience/security review and paper-retirement criteria.'],
    ['Support and optimisation', 'Month 10 onward', 'Health checks, incident review, adoption support, change requests and capacity/cost review.'],
], [1.65, 1.15, 4.2], 7.45)
heading('14.3 Pilot and rollout gates', 2)
table(['Gate', 'Minimum evidence'], [
    ['Pilot entry', 'Training completed; devices enrolled; KYC and SMS paths tested; ledger/reconciliation rehearsal balances; provider outage playbook; branch owner named.'],
    ['Pilot exit', 'At least 2-4 weeks of stable operations; daily cash lifecycle reconciled; no unresolved severity-1 ledger defect; member receipt and support metrics reviewed.'],
    ['Each branch wave', 'Opening balances signed; staff trained; device inventory complete; support rota active; paper parallel period defined; exit reconciliation complete.'],
    ['Credit launch', 'Client-approved policy and limits; accountant mapping; bureau path certified; PAR reports reconciled; adverse/dispute process active.'],
], [1.5, 5.5], 8.0)

heading('15 Bottlenecks failure modes and vendor lock-in', 1)
heading('15.1 Likely bottlenecks', 2)
table(['Bottleneck', 'Why it matters', 'Mitigation / owner'], [
    ['NIA/IVSP approval, contract or uptime', 'Can delay compliant onboarding or create exception backlog.', 'Start due diligence in Phase 0; adapter isolation; exception queue and service metrics. Owner: compliance + technical lead.'],
    ['USSD short-code lead time and network coverage', 'Member access channel may lag the app rollout.', 'Apply in Phase 0; support shared-code contingency; contract-test provider adapters. Owner: commercial lead.'],
    ['Source data quality', 'Incorrect opening balances damage trust immediately.', 'Branch-signed migration, sampled source review and member confirmation. Owner: finance + branch management.'],
    ['Offline operations and phone logistics', 'A sound app fails if devices are unmanaged, uncharged or untrained.', 'MDM, spare-device pool, data/power plan, field champions and queue-age monitoring. Owner: operations.'],
    ['Reconciliation discipline', 'Technology cannot compensate for missing cash counts/deposits.', 'Named close owners, variance SLA, management review and no-scale gate. Owner: finance.'],
    ['Rule churn', 'Changing products/limits during implementation creates unsafe ambiguity.', 'Phase 0 rule book and controlled change process. Owner: client sponsor.'],
], [1.7, 2.8, 2.5], 7.75)
heading('15.2 Failure-mode playbooks', 2)
table(['Failure', 'Safe system behaviour', 'Recovery evidence'], [
    ['Agent app offline for days', 'Queue immutable commands; show pending status and local cash summary; do not infer central balance changes.', 'Sync reconciliation report, queue hashes and supervisor review.'],
    ['Duplicate USSD/webhook request', 'Reject/return original result by idempotency key; no second posting.', 'Callback signature, request hash and original posting correlation.'],
    ['KYC provider unavailable', 'Keep intake as draft/retry-required; no verified status by proxy.', 'Provider incident record, retry attempts and compliance queue.'],
    ['Database/service outage', 'Fail closed for financial commands; read-only degraded views only if clearly stale-labeled; preserve local queues.', 'PITR/revision recovery, journal integrity check and incident review.'],
    ['Wrong amount or cash variance', 'Open a case; enforce controlled reversal/correction rather than edit.', 'Original event, counted evidence, approval and linked journal batches.'],
    ['Lost agent phone', 'MDM lock/wipe/revoke; prohibit new commands; replace through re-enrolment.', 'Device revoke event, last sync time, custody review and replacement record.'],
], [1.6, 3.7, 1.7], 7.5)
heading('15.3 Vendor lock-in mitigation', 2)
bullets([
    'Keep PostgreSQL schema, migrations, ledger rules and API contracts in Stormglide-controlled source repositories; test regular portable database restore.',
    'Use provider adapter interfaces for NIA, SMS, USSD, bureau, MDM and future payments. Store business IDs rather than provider-only identifiers where possible.',
    'Use Firebase for managed app services but keep core domain logic in Cloud Run and PostgreSQL. Export Firestore/Storage/config data on an exercised schedule.',
    'Avoid writing business decisions into provider workflows, Firebase Rules alone or unversioned admin screens. Product rules belong in versioned database configuration with approval history.'
])

heading('16 Open decisions and compliance assumptions', 1)
table(['Decision / assumption', 'Why it blocks or shapes build', 'Owner / required evidence'], [
    ['Legal entity and Bank of Ghana licence category', 'Determines reports, permissible products, supervisory expectations and may affect digital request scope.', 'Client counsel / regulator confirmation.'],
    ['NIA/IVSP provider and permitted biometric path', 'Sets onboarding hardware, consent wording, fallback and commercial cost.', 'Client/compliance + contracted provider certification.'],
    ['Restricted KYC policy', 'Must define what can be opened, credited/debited, notified and for how long when verification fails.', 'Compliance-approved policy; NIA/BoG guidance validation.'],
    ['Products and accounting rules', 'Rates, fees, commissions, catch-up, withdrawal reservations, liens, approvals and write-offs must be deterministic.', 'Signed Business Rules Specification and accountant chart-of-accounts mapping.'],
    ['Branch/cash process', 'Defines custody states, cash limits, close cut-off, deposit timing and variance escalation.', 'Finance/operations SOP and pilot evidence.'],
    ['Privacy and retention policy', 'Controls KYC image, card data, staff access logs and deletion/legal hold rules.', 'Data Protection Officer/legal review.'],
    ['Languages and accessibility', 'Affects USSD labels, training, member support and translations.', 'Client sponsor and field research.'],
    ['Provider selections and SLAs', 'Cost, failover and availability planning depend on real contract terms.', 'Procurement/commercial comparison and signed contracts.'],
], [1.75, 3.15, 2.15], 7.5)

heading('17 Technical cost and complexity factors for later pricing', 1)
p('No client price is included here. The following factors should feed the later Phase 0 / Phase 1 estimate and the subsequent module estimates. They are separated from pricing so commercial decisions do not distort safe technical scope.')
table(['Factor', 'Why it changes effort or running cost'], [
    ['Number of branches, agents, towns and current members', 'Sets device enrolment, training, migration volume, operational support and scale testing.'],
    ['KYC/IVSP provider and biometric hardware', 'Changes integration certification, per-check fees, field hardware, exception workflow and privacy controls.'],
    ['Existing data quality', 'Impacts mapping, reconciliation, duplicate resolution, opening-balance validation and parallel-run duration.'],
    ['USSD/SMS provider terms', 'Affects short-code setup, per-session/SMS costs, sandbox availability, delivery reporting and fallback adapters.'],
    ['Financial product complexity', 'Multiple products, tiered fees, commissions, liens, guarantees, supplier flows and exceptions expand ledger/posting and test matrix.'],
    ['Compliance/reporting obligations', 'Determines audit retention, bureau exports, report templates, approval evidence, security review and legal validation.'],
    ['Availability and disaster recovery target', 'HA database topology, backup frequency, cross-region recovery, monitoring and support coverage scale with RPO/RTO.'],
    ['Managed device programme', 'Phones, MDM licences, data bundles, replacement pool, private distribution, power and field support are operating costs.'],
    ['Migration and rollout model', 'Per-branch cleansing, training, on-site support, paper parallel-run and opening-balance confirmation drive delivery effort.'],
    ['Future payment movement', 'Mobile Money collections/payouts add regulated integrations, settlement/reversal reconciliation, limits, fraud monitoring and provider fees; estimate separately.'],
], [2.2, 4.85], 8.0)

heading('18 Risk register', 1)
table(['ID', 'Risk', 'Likelihood', 'Impact', 'Mitigation and trigger'], [
    ['R1', 'Incorrect opening balances or duplicate members.', 'High', 'High', 'Staged migration, signed reconciliation, member confirmation. Trigger: branch totals fail validation.'],
    ['R2', 'Offline sync posts duplicate/lost collections.', 'Medium', 'Critical', 'Idempotency, queue persistence, invariant tests, pilot reconciliation. Trigger: unmatched client/server command.'],
    ['R3', 'KYC provider/biometric process fails in the field.', 'Medium', 'High', 'Approved hardware, retry ladder, restricted compliance case. Trigger: no-match/error threshold.'],
    ['R4', 'Agent/device compromise or loss.', 'Medium', 'High', 'MDM, encryption, device binding, revoke/recovery SOP. Trigger: device posture or location anomaly.'],
    ['R5', 'Cash custody mismatch is hidden or resolved informally.', 'Medium', 'Critical', 'Mandatory close/reconciliation, immutable cases, manager/audit alerts. Trigger: handover/close/deposit variance.'],
    ['R6', 'Business rules change after build starts.', 'High', 'High', 'Signed Phase 0 rules and change control. Trigger: unapproved rule request.'],
    ['R7', 'USSD or SMS provider outage/delay.', 'Medium', 'Medium', 'Adapters, queue/retry, channel messaging and commercial SLA. Trigger: callback/receipt SLA breach.'],
    ['R8', 'Privacy/regulatory requirement misunderstood.', 'Medium', 'Critical', 'Counsel/DPO/compliance review, documented assumptions and pre-production assessment. Trigger: unresolved policy question.'],
    ['R9', 'Rollout exceeds operations capacity.', 'High', 'High', 'Wave gates, branch champions, hardware staging and parallel run. Trigger: training/device readiness below threshold.'],
    ['R10', 'Firebase or cloud cost grows unexpectedly.', 'Medium', 'Medium', 'Budgets/alerts, data lifecycle, cost dashboard, load tests and monthly review. Trigger: forecast exceeds approved run-rate.'],
], [0.45, 1.75, 0.65, 0.65, 3.7], 7.15)

heading('Conclusion and next internal actions', 1)
p('The system is buildable, but it should be built as a controlled financial operating system rather than a collection of apps. The technical centre is the PostgreSQL-backed Financial Core and append-only ledger; Firebase accelerates secure, managed delivery around that centre. Before implementation, Phase 0 must produce an approved business-rules specification, KYC/IVSP provider plan, accounting map, data-migration evidence, device operations plan and pilot acceptance gates.')
num([
    'Review and sign the Phase 0 decision register with client sponsor, finance, compliance, operations and legal owners.',
    'Run provider due diligence for NIA/IVSP, USSD, SMS, MDM and database hosting, including sandbox/certification and data-processing terms.',
    'Create the detailed data model, OpenAPI command contracts, posting-rule catalogue and screen/field workflow prototypes for the Core Susu pilot.',
    'Estimate and price only after the open decisions, migration sample and operating model are evidenced.'
])

# Set metadata
props = doc.core_properties
props.title = 'Microfinance and Susu Internal Technical Blueprint'
props.subject = 'Internal build architecture, controls and delivery plan'
props.author = 'STORMGLIDE Enterprise Systems'
props.comments = 'Derived from supplied project blueprint and planning decisions.'

doc.save(OUT)
print(OUT)
