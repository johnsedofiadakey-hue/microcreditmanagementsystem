from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path('/Users/truth/Developer/Microfinance Susu Management System')
OUT = ROOT / 'outputs/Stormglide_Microfinance_Susu_Commercial_Proposal_and_Invoice.docx'
LOGO = Path('/Users/truth/Downloads/logo-image copy.png')

NAVY = '071B37'
BLUE = '006BFF'
PALE_BLUE = 'EEF5FF'
PALE_GRAY = 'F4F6F8'
MID = '5C6A79'
TEXT = '17212B'
WHITE = 'FFFFFF'
GREEN = 'EAF5ED'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def border(cell, color='D7DEE6'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right'):
        tag = 'w:' + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag); borders.append(node)
        node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'6'); node.set(qn('w:color'),color)

def padding(cell, top=90, start=110, bottom=90, end=110):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side); mar.append(node)
        node.set(qn('w:w'),str(value)); node.set(qn('w:type'),'dxa')

def repeat(row):
    trPr = row._tr.get_or_add_trPr(); node = OxmlElement('w:tblHeader'); node.set(qn('w:val'),'true'); trPr.append(node)

def dont_split(row):
    trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))

def set_font(run, size=9.4, color=TEXT, bold=False, italic=False):
    run.font.name = 'Aptos'; run._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold; run.italic = italic

def ptext(doc, text='', style='Body', align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if text: set_font(p.add_run(text), 9.6)
    return p

def title(doc, text):
    p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(text); return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next = True; p.add_run(text); return p

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def add_logo(doc, width=Inches(4.6)):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(22); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(); r.add_picture(str(LOGO), width=width)
    doc.inline_shapes[-1]._inline.docPr.set('descr','Stormglide logo')

def data_table(doc, headers, rows, widths, money_cols=()):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    hdr = t.rows[0]; repeat(hdr); dont_split(hdr)
    for i, label in enumerate(headers):
        c = hdr.cells[i]; c.width = Inches(widths[i]); shade(c,NAVY); border(c,NAVY); padding(c)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run(label); set_font(r,8.6,WHITE,True)
        if i in money_cols: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for rix,row in enumerate(rows):
        rr = t.add_row(); dont_split(rr)
        for i,val in enumerate(row):
            c=rr.cells[i]; c.width=Inches(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(c, PALE_BLUE if rix%2==0 else WHITE); border(c); padding(c)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
            if i in money_cols: p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            r=p.add_run(str(val)); set_font(r,8.45,TEXT, i==0 and len(headers)==2)
    spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(2)
    return t

def bullet_list(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(item)

doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Inches(8.27),Inches(11.69)
sec.top_margin,sec.bottom_margin=Inches(.68),Inches(.63)
sec.left_margin,sec.right_margin=Inches(.70),Inches(.70)

styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.6); styles['Normal'].font.color.rgb=RGBColor.from_string(TEXT)
body=styles.add_style('Body',WD_STYLE_TYPE.PARAGRAPH); body.base_style=styles['Normal']; body.paragraph_format.space_after=Pt(6); body.paragraph_format.line_spacing=1.10
for name,size,color,before,after in [
    ('Title',25,NAVY,0,7),('Subtitle',12.3,MID,0,12),('Heading 1',16.5,NAVY,15,6),('Heading 2',11.8,BLUE,9,4),('Heading 3',10.2,NAVY,7,3)]:
    s=styles[name]; s.font.name='Aptos Display' if name in ('Title','Heading 1') else 'Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=name!='Subtitle'; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    pPr=s.element.pPr
    if pPr is not None:
        pb=pPr.find(qn('w:pBdr'))
        if pb is not None: pPr.remove(pb)
for sname in ('List Bullet','List Number'):
    styles[sname].font.name='Aptos'; styles[sname].font.size=Pt(9.2)

header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('STORMGLIDE'); set_font(r,8,NAVY,True)
r=header.add_run('.io'); set_font(r,8,BLUE,True)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Stormglide | stormglide.io | john@stormglide.io | 0547738678'); set_font(r,7.4,MID)

# Cover
add_logo(doc)
p=ptext(doc,'COMMERCIAL PROPOSAL AND PRO FORMA INVOICE',align=WD_ALIGN_PARAGRAPH.CENTER,before=0,after=7)
for run in p.runs: set_font(run,9.8,BLUE,True)
title(doc,'Microfinance and Susu System')
p=ptext(doc,'Full platform development and managed service investment',style='Subtitle',align=WD_ALIGN_PARAGRAPH.CENTER)

meta=data_table(doc,['Prepared by','Project commercial details'],[
    ('Stormglide','Client: To be confirmed'),
    ('Quote reference: STG MFS 001','Issued: 29 September 2026'),
    ('Currency: Ghana cedi','Delivery basis: staged implementation')
],[2.75,3.95])

heading(doc,'Project investment',1)
ptext(doc,'This proposal covers the full target platform described in the client execution guide. The platform is delivered in controlled stages so that the business proves the core Susu operation before larger credit, hire purchase and welfare workflows are activated.')
price=data_table(doc,['Commercial item','Amount'],[
    ('One time full platform development fee','GH₵55,000'),
    ('Monthly managed service after initial live pilot','GH₵8,500 per month'),
],[4.6,2.1],money_cols=(1,))
ptext(doc,'The development fee is the agreed professional fee for the full staged scope below. Third-party provider charges, statutory charges and optional devices are not included in these figures.',before=2)

heading(doc,'What the development fee includes',1)
ptext(doc,'The GH₵55,000 development investment covers product design, engineering, integration work, testing, launch preparation, initial training support and project documentation for the following staged platform package.')
page_break(doc)

heading(doc,'Development scope and fee schedule',1)
ptext(doc,'The line items below explain how the one-time development investment is allocated. They are delivered as connected parts of one platform, not as unrelated standalone products.')
data_table(doc,['Platform component','Included work','Development fee'],[
    ('Shared platform foundation','Role-based access, member records, central financial controls, audit trail, reporting base, security configuration and environment setup.','GH₵10,000'),
    ('Admin and branch portal','Head-office, branch, town-manager, cashier and supervisor dashboards; work queues, approvals, cash close and reporting.','GH₵10,000'),
    ('Android field agent app','Paired-agent workflow, offline collection, member lookup, receipts, route activity, device controls and cash handover.','GH₵9,000'),
    ('Susu savings operations','Collection cycles, savings records, handovers, daily close, reconciliation, exceptions and member receipts.','GH₵8,000'),
    ('Member PWA USSD and SMS','Member-facing PWA, basic USSD checking and requests, SMS receipts, reminders and service updates.','GH₵5,000'),
    ('Microcredit system','Loan request, assessment, approval, repayment tracking, arrears workflow and reporting.','GH₵5,000'),
    ('Hire purchase system','Catalogue, customer application, approvals, supplier and delivery tracking, instalment plans and arrears view.','GH₵4,000'),
    ('Welfare system','Membership rules, contributions, claim requests, approvals, payment tracking and case history.','GH₵2,000'),
    ('Public website and launch','Company website, product information, branch contacts, enquiry journey, FAQs and launch configuration.','GH₵2,000'),
    ('Total one-time development investment','','GH₵55,000'),
],[1.52,3.70,1.48],money_cols=(2,))

heading(doc,'Managed service after launch',1)
ptext(doc,'The monthly managed service begins after the initial live pilot is launched. It keeps the live platform stable, monitored and supported while the company operates it and progresses through the agreed stages.')
data_table(doc,['Included in GH₵8,500 per month','Service outcome'],[
    ('Platform monitoring and incident response','Operational issues are monitored, triaged and addressed through an agreed support process.'),
    ('Security and dependency maintenance','Routine security updates, access reviews and platform health maintenance.'),
    ('Backup oversight and recovery readiness','Backups and recovery procedures are checked so important operating data can be restored when required.'),
    ('Minor fixes and operational support','Support for defects, configuration questions, user access and practical day-to-day operations.'),
    ('Performance and service reporting','Regular view of service health, important incidents and actions requiring business input.')
],[3.25,3.45])
page_break(doc)

heading(doc,'Delivery and payment milestones',1)
ptext(doc,'The full project is released through stages. This payment schedule keeps delivery, testing and the company\'s operational readiness aligned.')
data_table(doc,['Milestone','What is achieved','Amount'],[
    ('1. Project commencement','Discovery, operating rules, product setup, architecture, design and pilot preparation begin.','GH₵22,000'),
    ('2. Core pilot ready','Core Susu workflow, admin and branch portal, agent app, collection, handover and basic reporting are ready for pilot acceptance.','GH₵13,750'),
    ('3. Member access and controlled expansion','Member PWA, USSD and SMS service access, exception controls and controlled rollout capabilities are completed.','GH₵11,000'),
    ('4. Final module delivery and handover','Credit, hire purchase, welfare, final hardening, documentation and handover are completed.','GH₵8,250'),
    ('Total development fee','','GH₵55,000'),
],[1.65,3.72,1.33],money_cols=(2,))

heading(doc,'Service and provider cost boundaries',1)
ptext(doc,'Stormglide will build and manage the platform described in this proposal. The following costs remain direct third-party or client operating costs because they depend on the company\'s chosen providers, usage and regulatory arrangements.')
bullet_list(doc,[
    'SMS messages, USSD short-code or session charges, messaging-provider fees and telecom usage.',
    'NIA IVSP, biometric verification, identity-provider, credit-bureau or related data-provider fees where used.',
    'Cloud hosting, Firebase, Google Cloud, database, storage, domain, email, app-store and monitoring charges.',
    'Company phones, supported biometric devices, staff data bundles, branch internet, banking and physical cash operations.',
    'Mobile-money payment integration, payment transaction fees and any separately approved future payment-collection scope.',
    'Taxes, levies and regulatory fees that may apply to the client or its external providers.'
])
ptext(doc,'Major new features, new product lines, material integrations or scope changes outside this proposal are quoted separately before work begins.',before=3)

heading(doc,'Assumptions and next step',1)
ptext(doc,'The client will nominate a pilot branch, confirm the relevant operating rules, provide timely access to approved providers and make the right staff available for testing and training. Delivery dates are confirmed after those dependencies and the pilot plan are agreed.')
ptext(doc,'Acceptance of this proposal confirms the development investment and the managed-service arrangement described here. Final commercial terms will be captured in the project agreement before implementation starts.')

sign=data_table(doc,['For the client','For Stormglide'],[
    ('Name: ______________________________\nTitle: ______________________________\nSignature: __________________________\nDate: ______________________________','Name: John Dakey\nCompany: Stormglide\nEmail: john@stormglide.io\nContact: 0547738678')
],[3.35,3.35])

doc.core_properties.title='Microfinance and Susu System Commercial Proposal and Pro Forma Invoice'
doc.core_properties.subject='Commercial proposal and investment summary'
doc.core_properties.author='Stormglide'
doc.core_properties.comments=''
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
