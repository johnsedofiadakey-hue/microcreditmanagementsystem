from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path('/Users/truth/Developer/Microfinance Susu Management System/outputs/Microfinance_Susu_Client_Experience_and_Execution_Guide.docx')
LOGO = Path('/Users/truth/Downloads/logo-image copy.png')
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = '071B37'
BLUE = '006BFF'
PALE_BLUE = 'EEF5FF'
PALE_GRAY = 'F2F4F6'
MID_GRAY = '5A6673'
WHITE = 'FFFFFF'
GREEN = 'E7F3EA'
AMBER = 'FFF3D6'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D4DCE3'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:color'), color)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

def set_no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit'); trPr.append(el)

def cell_para(cell, text='', bold=False, color='17212B', size=8.8, align=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align is not None: p.alignment = align
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    return p

def add_bullet(cell, text, size=8.5):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(.10)
    p.paragraph_format.first_line_indent = Inches(-.10)
    r = p.add_run('• ' + text); r.font.size = Pt(size)
    return p

def remove_title_border(style):
    pPr = style.element.pPr
    if pPr is None: return
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is not None: pPr.remove(pBdr)

def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def add_logo(paragraph, width=Inches(4.7)):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(LOGO), width=width)
    shape = paragraph.part.document.inline_shapes[-1]
    shape._inline.docPr.set('descr', 'Stormglide logo')
    return shape

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p

def add_text(doc, text, bold_lead=None):
    p = doc.add_paragraph(style='Body')
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def table(doc, headers, rows, widths=None, header_color=NAVY, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    set_repeat_header(hdr); set_no_split(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]; set_cell_shading(c, header_color); set_cell_border(c, header_color); set_cell_margins(c)
        cell_para(c, h, True, WHITE, font_size)
        if widths: c.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        tr = t.add_row(); set_no_split(tr)
        for i, value in enumerate(row):
            c = tr.cells[i]; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_shading(c, PALE_BLUE if ridx % 2 == 0 else WHITE); set_cell_border(c); set_cell_margins(c)
            if widths: c.width = Inches(widths[i])
            if isinstance(value, list):
                cell_para(c, value[0] if value else '', False, '17212B', font_size)
                for bullet in value[1:]: add_bullet(c, bullet, font_size-0.2)
            else: cell_para(c, str(value), False, '17212B', font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.70), Inches(.70)

styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(9.6); normal.font.color.rgb = RGBColor.from_string('1E2933')
body = styles.add_style('Body', WD_STYLE_TYPE.PARAGRAPH); body.base_style = styles['Normal']; body.paragraph_format.space_after = Pt(6); body.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in [('Title', 27, NAVY, 0, 8), ('Subtitle', 13, MID_GRAY, 0, 16), ('Heading 1', 17, NAVY, 16, 7), ('Heading 2', 12.4, BLUE, 10, 4), ('Heading 3', 10.2, NAVY, 7, 3)]:
    st = styles[name]; st.font.name = 'Aptos Display' if name in ('Title','Heading 1') else 'Aptos'; st.font.size = Pt(size); st.font.bold = name not in ('Subtitle',); st.font.color.rgb = RGBColor.from_string(color); st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
remove_title_border(styles['Title'])
for name in ('List Bullet', 'List Number'):
    styles[name].font.name = 'Aptos'; styles[name].font.size = Pt(9.2)

# Header and footer
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run('STORMGLIDE'); r.font.name = 'Aptos'; r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
r = header.add_run('.io'); r.font.name = 'Aptos'; r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Prepared by Stormglide | stormglide.io | john@stormglide.io | 0547738678')
r.font.name = 'Aptos'; r.font.size = Pt(7.5); r.font.color.rgb = RGBColor.from_string(MID_GRAY)

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(35); p.paragraph_format.space_after = Pt(0)
add_logo(p)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(8)
r = p.add_run('PREPARED BY STORMGLIDE'); r.font.name = 'Aptos'; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Microfinance and Susu System')
p = doc.add_paragraph(style='Subtitle'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Client execution guide for a connected field, branch and member operation')

t = doc.add_table(rows=1, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
set_repeat_header(t.rows[0])
for i, (title, note) in enumerate([('HEAD OFFICE', 'Sees the whole business'), ('BRANCH AND TOWN', 'Runs the daily operation'), ('FIELD AGENTS', 'Serve members in the community'), ('MEMBERS', 'Receive clear updates and make requests')]):
    c = t.rows[0].cells[i]; c.width = Inches(1.67); set_cell_shading(c, PALE_BLUE); set_cell_border(c, 'B8CEDF'); set_cell_margins(c, 180, 90, 180, 90)
    cell_para(c, title, True, NAVY, 8.5, WD_ALIGN_PARAGRAPH.CENTER)
    p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.space_after = Pt(0); rr = p2.add_run(note); rr.font.size = Pt(8); rr.font.color.rgb = RGBColor.from_string(MID_GRAY)

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(24); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('A clear, practical view of how the platform will support the company, its teams and its members.'); r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string('40505F')
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(30); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('stormglide.io  |  john@stormglide.io  |  0547738678'); r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(NAVY)
add_page_break(doc)

add_heading(doc, '1. The simple picture', 1)
add_text(doc, 'This is one connected operating system for Susu savings first, with microcredit, hire purchase and welfare added in controlled stages. It gives every person the right view of the work they are responsible for, while keeping one central record of what happened.')
table(doc, ['Place', 'What it is for', 'What the person sees'], [
    ('Public website', 'Explain the company, products, branches and how to join.', 'Company information, branch contacts, package explanations, calculator, enquiry form and FAQs.'),
    ('Agent Android app', 'Help paired field agents serve members and record a collection even when the network is poor.', 'Today\'s route, member search, collection screen, offline queue, cash handover summary and request forms.'),
    ('Branch and head-office portal', 'Run branches, review requests, receive cash, close the day and oversee performance.', 'Dashboards, work queues, member records, cash controls, reports, cases and audit history.'),
    ('Member SMS and USSD', 'Give members a simple way to receive proof and check or request services.', 'Receipts, reminders, balances, recent activity and request options when released.')
], [1.32, 2.36, 2.85], font_size=8.5)
add_text(doc, 'The system is designed to record, control and explain the operation. In the initial release, cash is still collected and paid through the company\'s normal physical process. The platform does not automatically move money from a member\'s phone or bank account.')

add_heading(doc, '2. How the four sides communicate', 1)
add_text(doc, 'Each person uses a different doorway, but nobody has a separate set of books. All important activity becomes part of the same central record, with the correct people notified or given a task.')
flow = doc.add_table(rows=1, cols=4); flow.alignment = WD_TABLE_ALIGNMENT.CENTER; flow.autofit = False
set_repeat_header(flow.rows[0])
for i, (title, body, shade) in enumerate([('MEMBER', 'Pays, asks, checks and receives proof.', GREEN), ('AGENT PAIR', 'Records the visit and carries the collection to branch.', PALE_BLUE), ('BRANCH', 'Receives cash, reviews requests and closes the day.', AMBER), ('HEAD OFFICE', 'Monitors performance, risks, approvals and policy.', PALE_GRAY)]):
    c = flow.rows[0].cells[i]; c.width = Inches(1.67); set_cell_shading(c, shade); set_cell_border(c); set_cell_margins(c, 170, 100, 170, 100)
    cell_para(c, title, True, NAVY, 9, WD_ALIGN_PARAGRAPH.CENTER)
    pp = c.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.space_after = Pt(0); rr = pp.add_run(body); rr.font.size = Pt(8); rr.font.color.rgb = RGBColor.from_string('344454')
add_text(doc, 'Practical flow: Member payment -> agent pair record -> member receipt after sync -> branch cash handover -> daily branch close -> bank deposit match -> management oversight.', bold_lead='Practical flow: ')
add_bullets(doc, [
    'A member does not need a smartphone to be included. SMS and USSD support basic access.',
    'An agent can continue working during a poor network period; the app keeps a protected queue and sends the records when it reconnects.',
    'Branch and head-office users see the part of the operation that belongs to their role, not unrestricted access to every action.'
])

add_heading(doc, '3. A normal collection day', 1)
table(doc, ['Step', 'What happens', 'Who sees the result'], [
    ('1. Visit', 'An agent pair meets the member and takes the physical payment according to the agreed collection process.', 'Agent pair records the visit; member sees the amount stated back to them.'),
    ('2. Record', 'The agent selects the member, enters the amount and saves the collection. If offline, it waits safely in the phone until sync.', 'Agent pair sees a pending or synced status; branch sees it after sync.'),
    ('3. Proof', 'Once the record reaches the central system, the member receives an SMS receipt where messaging is available.', 'Member receives proof; supervisors can see delivery or sync exceptions.'),
    ('4. Handover', 'The agent pair hands cash to the branch cashier. Both sides record and confirm the handover.', 'Cashier, manager and head office see expected versus received cash.'),
    ('5. Daily close', 'The branch counts cash, handles approved payouts, explains differences and closes the day.', 'Branch manager and finance team see the close status and any exception.'),
    ('6. Bank match', 'The physical bank deposit is matched to the branch close. Differences remain visible until resolved.', 'Finance and management see a clear reconciliation trail.')
], [0.58, 3.65, 2.3], font_size=8.25)

add_heading(doc, '4. At-a-glance responsibilities', 1)
table(doc, ['Role', 'Main dashboard view', 'Main controls', 'What is deliberately blocked'], [
    ('Overall administrator', 'Whole-company collections, branch performance, major approvals, exceptions and trends.', 'Company rules, products, role assignments, oversight and high-level approvals.', 'Cannot quietly rewrite past money records or approve their own request.'),
    ('Branch or town manager', 'Today\'s collections, agent activity, handovers, pending requests, branch close and local exceptions.', 'Manage assigned people and zones; approve within their authority; assign follow-up.', 'Cannot view or control another branch without permission; cannot exceed approval limit.'),
    ('Branch cashier', 'Expected cash, agent handovers, approved payout queue, cash float and close checklist.', 'Confirm receipts and handovers; prepare approved payments; close cash count.', 'Cannot approve their own payout or change ledger history.'),
    ('Supervisor', 'Agent pairs, routes, targets, missed activity, sync health and field checks.', 'Assign routes, coach agents, investigate flagged visits and escalate issues.', 'Cannot change a member balance or make a payment.'),
    ('Field agent pair', 'Today\'s members, collection form, offline queue, cash summary and requests started.', 'Register within policy, record collection, issue request and submit handover.', 'Cannot delete a collection, pay out cash, approve loans or edit historical balances.'),
    ('Member', 'SMS proof; later, USSD or PWA balance, recent activity and requests.', 'Check, request, report a problem and update permitted contact details.', 'Cannot transfer money or approve their own request through the initial member channels.')
], [1.1, 2.0, 2.05, 1.55], font_size=7.8)

add_heading(doc, '5. What the overall administrator sees', 1)
add_text(doc, 'The overall administrator sees the business from above. This is not a screen for doing every transaction personally; it is a screen for knowing what is healthy, what needs approval and where action is needed.')
table(doc, ['Dashboard area', 'What it answers'], [
    ('Today at a glance', 'How much was collected, how many active agent pairs worked, what has been handed in, and which branches have not closed.'),
    ('Branch comparison', 'Which towns or branches are on target, behind target, holding unusual cash or generating many exceptions.'),
    ('Approval and exception queue', 'Which large withdrawals, loans, reversals, write-offs or policy exceptions require the right independent approval.'),
    ('Risk and service watch', 'Which members have unresolved identity issues, which devices have concerns, which cases are overdue and where receipts or syncs are failing.'),
    ('Reports and policy', 'Company performance, product settings, role assignments and exportable reports for management review.')
], [2.05, 4.65], font_size=8.5)
add_text(doc, 'The administrator can set policy and oversee people. They cannot make a past collection disappear, change a balance without a trace or self-approve an action that requires an independent checker.')

add_heading(doc, '6. What a branch or town manager sees', 1)
add_text(doc, 'A branch or town manager has a practical daily view of their assigned operation. They see where field activity, cash and member requests need attention before they become problems.')
table(doc, ['On the manager dashboard', 'How it is used'], [
    ('Daily activity board', 'See expected collections, active agent pairs, visits not yet synced and members needing a follow-up.'),
    ('Cash and handover board', 'Compare collections recorded, cash handed in, cash received and the branch close position.'),
    ('Work queue', 'Review pending withdrawals, new loan or hire purchase requests, member updates and cases assigned to the branch.'),
    ('People and coverage', 'See teams, zones, route performance, training needs and absent or inactive devices.'),
    ('Branch health', 'See unresolved cash differences, delayed bank matches, complaints, KYC exceptions and overdue actions.')
], [2.1, 4.6], font_size=8.5)
add_text(doc, 'A manager can act only within an agreed branch and approval limit. The portal prevents self-approval, makes major actions require a second person and keeps a history of every decision.')

add_heading(doc, '7. What the field team sees and controls', 1)
add_heading(doc, 'Supervisor view', 2)
add_text(doc, 'Supervisors focus on delivery in the field: which pairs are active, which routes have been missed, whether phones are syncing and where a spot check or coaching visit is required. They can guide and escalate. They do not alter balances or complete a cash payout.')
add_heading(doc, 'Agent pair view', 2)
table(doc, ['Screen', 'What the agent can do', 'Important guardrail'], [
    ('Today\'s work', 'Open assigned route, see member list and view a simple collection target.', 'Only assigned members and approved work appear.'),
    ('Member search and profile', 'Find a member, confirm key details and see the information needed for service.', 'Sensitive data is limited to the task; full company records are not exposed.'),
    ('Collection', 'Enter payment, save a receipt reference and continue when the network is unavailable.', 'A saved collection cannot be deleted or silently changed by the agent.'),
    ('Offline queue', 'See items waiting to sync and resolve clear prompts after reconnecting.', 'The system detects duplicate or conflicting records before they enter the central record.'),
    ('Cash handover', 'Review expected cash and start a handover to the assigned branch cashier.', 'Both parties must confirm the handover; differences create an exception.'),
    ('Requests', 'Start a withdrawal, loan, hire purchase or welfare request for a member where that module is active.', 'Starting a request is not approval and does not release cash.')
], [1.45, 3.25, 1.95], font_size=8.1)
add_text(doc, 'Agent access is limited to approved company devices. Phone fingerprint or face unlock can protect the agent\'s app access, but it is not a customer Ghana Card fingerprint reader.')

add_heading(doc, '8. What the customer sees and controls', 1)
add_text(doc, 'The customer experience is intentionally simple. A member should understand what they paid, what was recorded and what is happening with any request without needing to visit a branch just to ask a basic question.')
table(doc, ['Member touchpoint', 'What the member sees or can do', 'What does not happen automatically'], [
    ('SMS receipt', 'Receive a receipt after a collection syncs, plus useful reminders and status updates.', 'A text message does not move money or approve a request.'),
    ('USSD', 'Use a basic phone and PIN to check permitted information, see recent activity and start simple requests when released.', 'The initial USSD service is for checking and requesting, not sending or withdrawing money.'),
    ('Member PWA', 'For smartphone members, view balance, activity, requests, notices and support options when released.', 'It does not bypass branch checks, identity controls or approval rules.'),
    ('Support and complaints', 'Report a missing receipt, disputed entry, lost phone or service issue through a defined support path.', 'A complaint does not delete a record; it opens a tracked case for review.')
], [1.35, 3.55, 1.75], font_size=8.15)
add_text(doc, 'The exact member screens are released in stages. The first practical promise is clear proof of collections and a controlled path to ask for help.')

add_heading(doc, '9. Identity, trust and a safe fallback', 1)
add_text(doc, 'The Ghana Card is the primary identity source for member onboarding. Where provider approval and integration are in place, the process uses the appropriate NIA IVSP route, such as face verification or a supported biometric device. A normal Android phone fingerprint reader protects the agent\'s device; it does not provide a customer biometric template for national verification.')
table(doc, ['If verification is not successful', 'What happens next'], [
    ('1. Quality check', 'The user is guided to retry with better lighting, camera position or a clearer capture.'),
    ('2. Approved alternative', 'Use a supported fingerprint or other approved IVSP route where available.'),
    ('3. Escalation', 'The case is sent through the approved NIA or IVSP escalation process.'),
    ('4. Controlled exception', 'The member is marked restricted and reviewed by compliance. No staff member gets an unrestricted override button.'),
    ('NHIS information', 'A photo and details may be captured only as supporting information. NHIS is not the primary KYC basis.')
], [2.05, 4.65], font_size=8.45)

add_heading(doc, '10. Requests, approvals and disputes', 1)
add_text(doc, 'Money-out actions and important changes should never depend on one person pressing one button. The system uses a clear maker-checker pattern: one authorised person starts or prepares an action, and a different authorised person checks and approves it.')
table(doc, ['Situation', 'How the system responds'], [
    ('Withdrawal or payout request', 'Member or staff starts the request; the correct branch role reviews it; an independent authorised role approves; the cashier pays only after approval; the record is finalised with proof.'),
    ('Loan, hire purchase or welfare request', 'The request is captured, assessed by the responsible team, approved within policy and issued only after the required checks are complete.'),
    ('Mistake or disputed collection', 'The case is logged, evidence is reviewed, and any correction follows an authorised reversal process. The original record remains visible in history.'),
    ('Cash difference', 'The branch close cannot hide it. It stays open as an exception, is assigned to an owner and is escalated if not resolved.'),
    ('Lost or replaced phone', 'The device is removed or re-bound through a controlled process before the user can continue.')
], [2.0, 4.7], font_size=8.35)

add_heading(doc, '11. What arrives first, and what follows', 1)
add_text(doc, 'The full vision is deliberately delivered in stages. This gives the company time to test real field use, train people, correct processes and build confidence before more complex products are switched on.')
table(doc, ['Delivery stage', 'What is introduced', 'What remains for a later stage'], [
    ('Stage 0 - preparation', 'Rules, roles, products, branch setup, forms, data preparation, provider readiness and pilot training.', 'No live company-wide rollout until operating rules are confirmed.'),
    ('Stage 1 - core Susu pilot', 'Branch portal, controlled agent app, member registration, core collection, SMS receipts, handover, daily close and basic management view.', 'Full credit, hire purchase, welfare and broad member self-service remain off until the pilot is accepted.'),
    ('Stage 2 - controlled expansion', 'More branches and agent pairs, member USSD access, stronger exception management, commission and broader reporting.', 'Expansion is paced by readiness, not just by calendar date.'),
    ('Stage 3 - credit and hire purchase', 'Loan workflows, credit assessment, hire purchase catalogue, approvals, repayment tracking and arrears handling.', 'Only introduced after core cash and reconciliation controls work reliably.'),
    ('Stage 4 - welfare and hardening', 'Welfare workflows, richer reports, operational refinements, resilience testing and wider rollout.', 'Future mobile-money payment integration is a separate scoped project, not an assumption in the first release.')
], [1.55, 3.5, 1.65], font_size=7.9)

add_heading(doc, '12. What this system will and will not do', 1)
table(doc, ['It will', 'It will not'], [
    (['Give each role a clear dashboard and task queue.', 'Create a central, traceable record of collections, handovers, approvals and cases.', 'Support field work through an offline-capable Android app.', 'Give members receipts and simple service access.'], ['Replace the company\'s cash handling policy or make physical cash disappear.', 'Allow staff to erase historic money activity without an audit trail.', 'Turn a request into a payout without the correct approval.', 'Treat an NHIS card photo as primary KYC or treat a phone fingerprint as customer national biometric verification.']),
], [3.35, 3.35], font_size=8.4)

add_heading(doc, '13. What we need to agree before build starts', 1)
add_bullets(doc, [
    'The pilot branch or town, the first agent pairs and the named operational owners.',
    'Susu product rules, collection cycles, withdrawals, approval limits, fees and exception handling.',
    'The physical cash handover, daily close and bank deposit process that the system will reflect.',
    'The identity verification provider path, company device policy and how restricted KYC cases are handled.',
    'Which later-stage modules are committed after the pilot: USSD, member PWA, loans, hire purchase and welfare.',
    'Training dates, a pilot acceptance checklist and the decision point for expanding to the next branch.'
])
add_text(doc, 'Commercial pricing, provider charges and the managed support arrangement are intentionally handled in the separate commercial proposal. This guide explains the operating picture and the scope boundaries that inform that proposal.')

add_heading(doc, '14. Closing picture', 1)
add_text(doc, 'The intended outcome is not simply a collection app. It is a disciplined operating picture: members get proof and a voice; agents get a practical field tool; branches get control of cash and work; and management gets visibility without compromising the checks that protect the business.')
add_text(doc, 'We begin with a working Susu pilot, prove the daily journey in one real branch, and grow from that foundation in controlled stages.', bold_lead='We begin with a working Susu pilot, ')

doc.core_properties.title = 'Microfinance and Susu System Client Execution Guide'
doc.core_properties.subject = 'Client-facing role and operating model guide'
doc.core_properties.author = 'Stormglide'
doc.core_properties.comments = ''
doc.save(OUT)
print(OUT)
